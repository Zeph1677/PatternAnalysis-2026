"""Build the visual COMP3710 2D-to-3D progress review."""
from pathlib import Path
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
EVIDENCE = ROOT / "evidence"
OUTPUT = EVIDENCE / "COMP3710_HipMRI_Hard_Feasibility_Review.docx"
OASIS_RUN = Path(r"D:\OneDrive - The University of Queensland\Desktop\Study\COMP3710_PatternRecognition\Lab\Demo2\Log\p4_job\p4_592041\Temp\unet")
HISTORY, PREDICTIONS = OASIS_RUN / "training_history.csv", OASIS_RUN / "test_predictions.png"
DSC, CURVE = OASIS_RUN / "test_dsc_per_class.csv", EVIDENCE / "oasis_2d_training_curve.png"

def shade(cell, color):
    node = OxmlElement("w:shd"); node.set(qn("w:fill"), color); cell._tc.get_or_add_tcPr().append(node)

def margins(section):
    section.top_margin = section.bottom_margin = Inches(0.42)
    section.left_margin = section.right_margin = Inches(0.55)

def heading(doc, text, size=11):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0, 0, 0)

def body(doc, text, size=9.2):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.03
    p.add_run(text).font.size = Pt(size)

def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(7.8); r.font.color.rgb = RGBColor(89, 89, 89)

def title(doc, main, sub):
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
    r = p.add_run(main); r.font.size = Pt(16); r.bold = True; r.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(5)
    r = p.add_run(sub); r.font.size = Pt(9.4); r.italic = True

def status_table(doc, rows):
    table = doc.add_table(rows=1, cols=3); table.autofit = False
    widths, headers = (Inches(1.55), Inches(1.05), Inches(4.35)), ("Stage", "Status", "Evidence")
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[i]; shade(cell, "1F4E79"); p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, values)):
            cell.width = widths[i]
            if row_index % 2: shade(cell, "F2F6FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(text).font.size = Pt(8.1)

def make_curve():
    history = np.genfromtxt(HISTORY, delimiter=",", names=True)
    row = np.genfromtxt(DSC, delimiter=",", names=True)
    fig, ax1 = plt.subplots(figsize=(8.2, 2.45), dpi=180)
    ax1.plot(history["epoch"], history["train_loss"], color="#C44E52", linewidth=1.8)
    ax1.set(xlabel="Epoch", ylabel="Training loss"); ax1.tick_params(axis="y", labelcolor="#C44E52"); ax1.grid(alpha=0.2)
    ax2 = ax1.twinx(); ax2.plot(history["epoch"], history["val_mean_dsc"], color="#1F77B4", linewidth=1.8)
    ax2.set_ylabel("Validation mean DSC"); ax2.tick_params(axis="y", labelcolor="#1F77B4"); ax2.set_ylim(0.94, 0.985)
    ax1.set_title("Completed OASIS 2D U-Net training"); fig.tight_layout(); fig.savefig(CURVE, bbox_inches="tight", facecolor="white"); plt.close(fig)
    return history, np.array([row[name].item() for name in row.dtype.names])

def main():
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    history, dsc = make_curve(); best = int(np.argmax(history["val_mean_dsc"]))
    doc = Document(); margins(doc.sections[0]); normal = doc.styles["Normal"]; normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    title(doc, "From 2D Segmentation to HipMRI 3D", "COMP3710 progress evidence and Hard project feasibility review")
    body(doc, "The 2D stage is complete and provides a verified segmentation pipeline, trained checkpoint and held-out visual evidence. The next stage changes dataset and clinical target: HipMRI volumetric prostate segmentation using matched plain and improved 3D U-Nets.")
    heading(doc, "Completed 2D result")
    table = doc.add_table(rows=1, cols=4)
    items = [("Dataset", "OASIS brain MRI slices"), ("Training", "80 epochs on A100"), ("Best validation DSC", f"{history['val_mean_dsc'][best]:.4f} at epoch {int(history['epoch'][best])}"), ("Test foreground mean DSC", f"{dsc[1:].mean():.4f}")]
    for cell, (label, value) in zip(table.rows[0].cells, items):
        shade(cell, "EAF2F8"); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}\n{value}"); r.font.size = Pt(8.1); r.bold = True
    doc.add_picture(str(CURVE), width=Inches(6.9)); caption(doc, "Figure 1. Training loss decreases while validation mean DSC stabilises near 0.977.")
    doc.add_picture(str(PREDICTIONS), width=Inches(4.65)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 2. Held-out OASIS examples: input MRI, ground truth and 2D U-Net prediction.")
    body(doc, f"Held-out per-class DSC: background {dsc[0]:.4f}; foreground classes 1-3 {dsc[1]:.4f}, {dsc[2]:.4f} and {dsc[3]:.4f}. All foreground classes exceed the earlier 2D target of 0.90. These are not HipMRI prostate results.", 8.5)
    section = doc.add_section(WD_SECTION.NEW_PAGE); margins(section)
    title(doc, "HipMRI 3D Stage", "Current implementation evidence and next experiment")
    heading(doc, "Why move from 2D to 3D")
    body(doc, "Independent 2D slices cannot directly enforce continuity across the prostate base, middle and apex. The 3D stage uses volumetric context and tests whether residual units plus deep supervision improve difficult boundaries enough to justify the additional memory and runtime.")
    status_table(doc, [("2D OASIS U-Net", "Complete", "80-epoch history, checkpoint, held-out DSC and prediction panel retained."), ("3D data pipeline", "Implemented", "Recursive NIfTI pairing; shape, affine, finite-intensity and integer-label validation."), ("Patient split", "Implemented", "Seed-42 train/validation/test manifest with zero-overlap checks."), ("3D models", "Implemented", "BaselineUNet3D and residual ImprovedUNet3D with two deep-supervision heads."), ("Automated verification", "8 tests pass", "Pairing, invalid-data rejection, patient split, forward/backward, loss and evaluation logic."), ("Real HipMRI training", "Next", "Data audit, Rangpur GPU smoke run, then matched baseline and improved full runs.")])
    heading(doc, "Planned fair comparison")
    flow = doc.add_table(rows=1, cols=5)
    for i, (cell, text) in enumerate(zip(flow.rows[0].cells, ["Audit HipMRI\nvolumes", "Freeze patient\nmanifest", "GPU smoke\nrun", "Train both\n3D models", "Held-out metrics\nand failure panels"])):
        shade(cell, "D9EAF7" if i < 2 else "EDEDED"); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(text); r.font.size = Pt(8.2); r.bold = True
    heading(doc, "Acceptance evidence for the 3D stage")
    body(doc, "Both models will use the same patient-disjoint manifest. Final evidence will report prostate Dice (target at least 0.70), per-class Dice/IoU, peak VRAM, per-case inference time and 3-5 held-out base/apex or boundary-leakage cases. Checkpoint selection uses validation data only.")
    heading(doc, "Resource risk and fallback")
    body(doc, "A short smoke run will determine patch size, batch size and mixed precision before full Slurm jobs. If the improved model does not reduce real held-out errors at acceptable cost, the plain 3D U-Net remains the fallback.")
    heading(doc, "Tutor confirmation requested")
    body(doc, "Please confirm the HipMRI label encoding, whether background is excluded from reported Dice, recommended patch and batch settings, compute budget, and required logs, checkpoints and failure panels.")
    footer = doc.sections[0].footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("COMP3710 progress review - 2D evidence is OASIS; 3D performance remains to be measured on HipMRI."); r.font.size = Pt(7.1); r.font.color.rgb = RGBColor(89, 89, 89)
    doc.save(OUTPUT); print(OUTPUT)

if __name__ == "__main__": main()
