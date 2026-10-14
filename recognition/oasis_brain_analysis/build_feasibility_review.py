"""Build a one-page COMP3710 feasibility review from verified local evidence."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
AUDIT = ROOT / "evidence" / "data_audit" / "dataset_manifest.json"
SMOKE = Path(r"E:\vscode\COMP3710_A1\artifacts\baseline\test_metrics.json")
OUTPUT = ROOT / "evidence" / "COMP3710_Midpoint_Feasibility_Review.docx"


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.2)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(10.2)
    run.font.color.rgb = RGBColor(31, 78, 121)


def _body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.size = Pt(9.4)


def _question(document: Document, number: int, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.first_line_indent = Inches(-0.08)
    paragraph.paragraph_format.line_spacing = 1.0
    number_run = paragraph.add_run(f"{number}. ")
    number_run.bold = True
    number_run.font.size = Pt(9.4)
    text_run = paragraph.add_run(text)
    text_run.font.size = Pt(9.4)


def main() -> None:
    audit = json.loads(AUDIT.read_text(encoding="utf-8"))
    smoke = json.loads(SMOKE.read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1)
    run = title.add_run("COMP3710 OASIS Brain Analysis Feasibility Review")
    run.bold = True
    run.font.size = Pt(15)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    subrun = subtitle.add_run("Current implemented module: baseline versus improved 2D U-Net tissue segmentation")
    subrun.italic = True
    subrun.font.size = Pt(10)

    table = document.add_table(rows=2, cols=3)
    table.autofit = False
    labels = [("Dataset audit", f"{audit['pair_count']:,} pairs; {audit['subject_count']} subjects"), ("Geometry", "256 x 256 for every pair"), ("Labels", "0/85/170/255 mapped to 0-3"), ("Split protocol", "Seed 42; subject-disjoint"), ("Smoke test", "Baseline, 1 epoch, 8 cases/split"), ("Consultation status", "Tutor check-off not yet recorded")]
    for cell, (label, value) in zip((cell for row in table.rows for cell in row.cells), labels):
        _shade(cell, "EAF2F8")
        _set_cell_text(cell, f"{label}\n{value}", bold=False)
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(2.47)

    _heading(document, "Purpose and decision requested")
    _body(document, "Research question: on the same subject-disjoint held-out set, can residual encoding, attention-gated skips and dropout improve class-wise tissue segmentation enough to justify their resource cost? I request confirmation that the Easy-difficulty scope and acceptance criteria are appropriate before full training.")

    _heading(document, "Method and acceptance criteria")
    _body(document, "Both models will use the same saved seed-42 subject-level manifest. The baseline is a plain 2D U-Net; the proposed model adds residual encoding, attention-gated skips and bottleneck dropout. Acceptance evidence: no subject overlap; four-class logits; class-wise held-out Dice/IoU; parameter, VRAM and inference records; and 3-5 ranked FP/FN failure panels. The course target is Dice >= 0.90 per tissue label.")

    _heading(document, "Verified local evidence")
    _body(document, f"Read-only audit of the local source found {audit['pair_count']:,} aligned image/mask pairs from {audit['subject_count']} subjects. Every pair is 256 x 256. The four observed raw codes (0, 85, 170 and 255) map unambiguously to classes 0-3. Four deterministic audit samples were saved. Unit tests cover nested pairing, subject allocation, encoded-label remapping, and the audit manifest.")

    _heading(document, "Smoke run: evidence of execution, not model quality")
    _body(document, f"A local GPU baseline smoke run recorded PyTorch {smoke['torch_version']} on {smoke['device']}: one epoch, 8 held-out cases, {smoke['trainable_parameters']:,} trainable parameters, {smoke['inference_seconds']:.3f} s inference, and {smoke['peak_vram_mb']:.1f} MB peak VRAM. Held-out mean Dice = {smoke['mean_dice']:.4f}; mean IoU = {smoke['mean_iou']:.4f}. These deliberately small-run values validate the pipeline and failure-reporting route only; they do not support a performance claim.")

    _heading(document, "Risks, budget and fallback")
    _body(document, "Risks: label imbalance, memory demand, thin tissue boundaries and a high mean masking class failures. Budget: train baseline and improved models on the identical full manifest; preserve checkpoints, learning curves, JSON metrics and 3-5 failure panels. Fallback: retain the simpler baseline if the improved model does not show a class-wise and qualitative benefit at acceptable resource cost.")

    _heading(document, "Questions for tutor confirmation")
    _question(document, 1, "Confirm the Easy scope and the four-class 2D OASIS target.")
    _question(document, 2, "Confirm Dice >= 0.90 per label as the evaluation target and whether background is reported separately.")
    _question(document, 3, "Confirm the required training budget and acceptable evidence for the midpoint record.")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Prepared from local audit and smoke artifacts. Tutor confirmation remains an external requirement.")
    footer_run.font.size = Pt(7.5)
    footer_run.font.color.rgb = RGBColor(89, 89, 89)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
